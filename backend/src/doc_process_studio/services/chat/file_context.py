from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi import UploadFile
from openpyxl import load_workbook
from pypdf import PdfReader

from ...models.conversation.file_context import (
    PreparedUploadedFile,
    UploadedFileContext,
)
from .attachments import load_uploaded_attachment_context, save_uploaded_attachment

TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".json",
    ".jsonl",
    ".yaml",
    ".yml",
    ".xml",
    ".html",
    ".htm",
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".java",
    ".go",
    ".rs",
    ".sql",
    ".log",
}
MAX_FILE_CHARACTERS = 1_000_000
SPECIAL_DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".xlsx"}


def decode_file_bytes(raw_bytes: bytes) -> str | None:
    for encoding in ("utf-8", "utf-8-sig", "gbk"):
        try:
            return raw_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue

    return None


def truncate_content(content: str) -> str:
    normalized_text = content.strip()
    if not normalized_text:
        return "文件内容为空，或当前无法提取有效文本。"

    if len(normalized_text) > MAX_FILE_CHARACTERS:
        return (
            normalized_text[:MAX_FILE_CHARACTERS]
            + "\n\n[文件内容过长，已截断后发送给模型]"
        )

    return normalized_text


def extract_pdf_text(raw_bytes: bytes) -> str:
    reader = PdfReader(BytesIO(raw_bytes))
    page_texts = [(page.extract_text() or "").strip() for page in reader.pages]
    return "\n\n".join(text for text in page_texts if text)


def extract_docx_text(raw_bytes: bytes) -> str:
    document = Document(BytesIO(raw_bytes))

    sections: list[str] = []
    paragraph_texts = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]
    if paragraph_texts:
        sections.append("\n".join(paragraph_texts))

    table_sections: list[str] = []
    for table_index, table in enumerate(document.tables, start=1):
        row_texts: list[str] = []
        for row in table.rows:
            cell_texts = [cell.text.strip() for cell in row.cells]
            non_empty_cells = [text for text in cell_texts if text]
            if non_empty_cells:
                row_texts.append(" | ".join(non_empty_cells))

        if row_texts:
            table_sections.append(f"[表格 {table_index}]\n" + "\n".join(row_texts))

    if table_sections:
        sections.append("\n\n".join(table_sections))

    return "\n\n".join(sections)


def extract_xlsx_text(raw_bytes: bytes) -> str:
    workbook = load_workbook(
        filename=BytesIO(raw_bytes),
        data_only=True,
        read_only=True,
    )

    sheet_sections: list[str] = []
    for sheet in workbook.worksheets:
        row_texts: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            cell_values = [
                str(cell).strip()
                for cell in row
                if cell is not None and str(cell).strip()
            ]
            if cell_values:
                row_texts.append(" | ".join(cell_values))

        if row_texts:
            sheet_sections.append(f"[工作表：{sheet.title}]\n" + "\n".join(row_texts))

    return "\n\n".join(sheet_sections)


def extract_special_document_text(
    suffix: str,
    raw_bytes: bytes,
) -> str | None:
    parser_map = {
        ".pdf": extract_pdf_text,
        ".docx": extract_docx_text,
        ".xlsx": extract_xlsx_text,
    }
    parser = parser_map.get(suffix)
    if not parser:
        return None

    return parser(raw_bytes)


async def extract_upload_file_context(
    upload_file: UploadFile,
) -> UploadedFileContext:
    raw_bytes = await upload_file.read()
    await upload_file.close()

    return _build_uploaded_file_context(
        filename=upload_file.filename or "未命名文件",
        content_type=upload_file.content_type,
        raw_bytes=raw_bytes,
    )


def _build_uploaded_file_context(
    *,
    filename: str,
    content_type: str | None,
    raw_bytes: bytes,
) -> UploadedFileContext:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix in SPECIAL_DOCUMENT_EXTENSIONS:
            extracted_text = extract_special_document_text(suffix, raw_bytes)
        else:
            extracted_text = decode_file_bytes(raw_bytes)
    except Exception:
        extracted_text = None

    if extracted_text is None and suffix not in TEXT_EXTENSIONS:
        content = (
            "当前后端无法提取该文档的有效文本内容，"
            "本次仅保留了文件名与类型信息。"
        )
    else:
        content = truncate_content(extracted_text or "")

    return UploadedFileContext(
        filename=filename,
        content_type=content_type,
        content=content,
    )


def _build_uploaded_file_section(file_context: UploadedFileContext) -> str:
    return "\n".join(
        [
            f"文件名：{file_context.filename}",
            f"类型：{file_context.content_type or 'unknown'}",
            "内容：",
            file_context.content,
        ]
    )


async def prepare_uploaded_files(
    *,
    upload_files: list[UploadFile],
    conversation_id: str,
    skill_id: str,
) -> tuple[list[PreparedUploadedFile], str | None]:
    if not upload_files:
        return [], None

    prepared_files: list[PreparedUploadedFile] = []
    for upload_file in upload_files:
        raw_bytes = await upload_file.read()
        await upload_file.close()
        file_context = _build_uploaded_file_context(
            filename=upload_file.filename or "未命名文件",
            content_type=upload_file.content_type,
            raw_bytes=raw_bytes,
        )
        attachment = save_uploaded_attachment(
            raw_bytes=raw_bytes,
            conversation_id=conversation_id,
            skill_id=skill_id,
            file_name=file_context.filename,
            mime_type=file_context.content_type,
            extracted_text=file_context.content,
        )
        prepared_files.append(
            PreparedUploadedFile(
                attachment=attachment,
                context=file_context,
            )
        )

    return prepared_files, (
        "以下是用户本次上传的文件内容，请你优先结合这些文件进行理解与回答：\n\n"
        + "\n\n---\n\n".join(
            _build_uploaded_file_section(prepared_file.context)
            for prepared_file in prepared_files
        )
    )

def build_persisted_uploaded_files_context(
    attachment_ids: list[str],
) -> str | None:
    normalized_attachment_ids: list[str] = []
    for attachment_id in attachment_ids:
        normalized_attachment_id = attachment_id.strip()
        if normalized_attachment_id and normalized_attachment_id not in normalized_attachment_ids:
            normalized_attachment_ids.append(normalized_attachment_id)

    if not normalized_attachment_ids:
        return None

    sections: list[str] = []
    for attachment_id in normalized_attachment_ids:
        metadata, extracted_text, is_expired = load_uploaded_attachment_context(attachment_id)
        if is_expired or metadata is None or not extracted_text:
            continue
        sections.append(
            "\n".join(
                [
                    f"文件名：{metadata.name}",
                    f"类型：{metadata.mime_type or 'unknown'}",
                    "内容：",
                    extracted_text,
                ]
            )
        )

    if not sections:
        return None

    return (
        "以下是当前会话中已绑定的历史上传文件内容，请你继续结合这些文件进行理解与回答：\n\n"
        + "\n\n---\n\n".join(sections)
    )
