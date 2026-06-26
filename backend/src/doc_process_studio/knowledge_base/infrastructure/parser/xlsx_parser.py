import logging
from dataclasses import dataclass
from io import BytesIO

from openpyxl import load_workbook

logger = logging.getLogger(__name__)


@dataclass
class ParsedSheet:
    sheet_name: str
    text: str


def parse_xlsx(file_bytes: bytes, file_name: str) -> list[ParsedSheet]:
    sheets: list[ParsedSheet] = []
    try:
        wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows: list[str] = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(cell) if cell is not None else "" for cell in row]
                line = " | ".join(cells)
                if line.strip("| "):
                    rows.append(line)
            text = "\n".join(rows).strip()
            if text:
                sheets.append(ParsedSheet(sheet_name=sheet_name, text=text))
        wb.close()
    except Exception:
        logger.warning("Failed to parse XLSX: %s", file_name, exc_info=True)
    return sheets
