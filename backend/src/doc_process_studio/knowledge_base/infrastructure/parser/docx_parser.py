import logging
from dataclasses import dataclass
from io import BytesIO

from docx import Document
from docx.document import Document as DocumentType

logger = logging.getLogger(__name__)


@dataclass
class ParsedSection:
    section_title: str
    text: str


def parse_docx(file_bytes: bytes, file_name: str) -> list[ParsedSection]:
    sections: list[ParsedSection] = []
    try:
        doc: DocumentType = Document(BytesIO(file_bytes))
        current_title = file_name
        current_paragraphs: list[str] = []

        for paragraph in doc.paragraphs:
            style_name = (paragraph.style.name or "").lower() if paragraph.style else ""
            is_heading = "heading" in style_name or style_name.startswith("toc")

            if is_heading and paragraph.text.strip():
                content = "\n".join(current_paragraphs).strip()
                if content:
                    sections.append(ParsedSection(section_title=current_title, text=content))
                current_title = paragraph.text.strip()
                current_paragraphs = []
                continue

            text = paragraph.text.strip()
            if text:
                current_paragraphs.append(text)

        trailing_content = "\n".join(current_paragraphs).strip()
        if trailing_content:
            sections.append(ParsedSection(section_title=current_title, text=trailing_content))

        if not sections:
            full_text = "\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())
            if full_text:
                sections.append(ParsedSection(section_title=file_name, text=full_text))
    except Exception:
        logger.warning("Failed to parse DOCX: %s", file_name, exc_info=True)
    return sections
