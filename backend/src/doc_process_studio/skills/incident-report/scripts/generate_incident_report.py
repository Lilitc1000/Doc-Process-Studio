#!/usr/bin/env python3
"""
DAS Fault Log Form Generator (Template-aligned)
使用参考 DOCX 模板直接填充，尽量保持与参考文档版式一致。
"""

from __future__ import annotations

import base64
import hashlib
import io
import json
import re
from datetime import datetime
from html import unescape
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Cm

SCRIPT_DIR = Path(__file__).resolve().parent
REFERENCE_TEMPLATE_PATH = (
    SCRIPT_DIR.parent
    / "references"
    / "DAS2 Fault Log Form Template.docx"
)

_IMAGE_DECODE_CACHE_MAX_ENTRIES = 48
_IMAGE_DECODE_CACHE_MAX_BYTES = 8 * 1024 * 1024
_IMAGE_DECODE_CACHE: dict[str, bytes] = {}
_IMAGE_DECODE_CACHE_ORDER: list[str] = []


def _to_text(value: Any, default: str = "N/A") -> str:
    if value is None:
        return default
    if isinstance(value, str):
        cleaned = value.strip()
        return cleaned or default
    if isinstance(value, (int, float, bool)):
        return str(value)
    if isinstance(value, list):
        items = [_to_text(item, "").strip() for item in value]
        items = [item for item in items if item]
        return "; ".join(items) if items else default
    if isinstance(value, dict):
        parts = []
        for key, item in value.items():
            item_text = _to_text(item, "").strip()
            if not item_text:
                continue
            key_text = str(key).replace("_", " ").strip()
            parts.append(f"{key_text}: {item_text}" if key_text else item_text)
        return "；".join(parts) if parts else default
    return str(value)


def _to_lines(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        lines = []
        for item in value:
            text = _to_text(item, "").strip()
            if text:
                lines.append(text)
        return lines
    text = _to_text(value, "").strip()
    if not text:
        return []
    if "\n" in text:
        rows = [line.strip("- ").strip() for line in text.splitlines()]
        return [row for row in rows if row]
    if ";" in text:
        rows = [line.strip() for line in text.split(";")]
        return [row for row in rows if row]
    return [text]


def _split_date_time(date_time_value: Any) -> tuple[str, str]:
    normalized = _to_text(date_time_value, "").strip()
    if not normalized:
        return "N/A", "N/A"
    parts = normalized.split()
    if len(parts) >= 2:
        return parts[0], parts[1]
    if ":" in normalized and "/" not in normalized:
        return "N/A", normalized
    return normalized, "N/A"


def _first_non_empty(*values: Any, default: str = "N/A") -> str:
    for value in values:
        text = _to_text(value, "").strip()
        if text:
            return text
    return default


def _normalize_event_sequence(value: Any) -> list[dict[str, str]]:
    if value is None:
        return []
    if not isinstance(value, list):
        return [{"time": "N/A", "event": _to_text(value), "evidence": "N/A"}]

    normalized = []
    for item in value:
        if isinstance(item, dict):
            normalized.append(
                {
                    "time": _to_text(item.get("time", "N/A")),
                    "event": _to_text(
                        item.get("event", item.get("description", item.get("detail", "N/A")))
                    ),
                    "evidence": _to_text(
                        item.get("evidence", item.get("source", item.get("proof", "N/A")))
                    ),
                }
            )
        else:
            normalized.append({"time": "N/A", "event": _to_text(item), "evidence": "N/A"})
    return normalized


def _normalize_actions(value: Any, action_type: str) -> list[dict[str, str]]:
    if value is None:
        return []
    if not isinstance(value, list):
        return [{"action": _to_text(value)}]

    normalized = []
    for item in value:
        if isinstance(item, dict):
            normalized.append(
                {
                    "action": _to_text(item.get("action", item.get("description", "")), ""),
                }
            )
        else:
            normalized.append({"action": _to_text(item, "")})
    return normalized


def _extract_action_lines(actions: Any) -> list[str]:
    lines: list[str] = []
    if isinstance(actions, list):
        for item in actions:
            if isinstance(item, dict):
                text = _to_text(item.get("action"), "").strip()
            else:
                text = _to_text(item, "").strip()
            if text:
                lines.append(text)
    elif isinstance(actions, dict):
        text = _to_text(actions.get("action"), "").strip()
        if text:
            lines.append(text)
    else:
        text = _to_text(actions, "").strip()
        if text:
            lines.append(text)
    return lines


def _build_reference_no() -> str:
    return f"DAS-{datetime.now().strftime('%Y%m%d')}-001"


STATUS_OPTION_FAULT_CLEARED = "fault_cleared"
STATUS_OPTION_TEMPORARILY_FIXED = "temporarily_fixed"
STATUS_OPTION_FOLLOW_UP_ACTION_REQUIRED = "follow_up_action_required"

SEVERITY_OPTION_NOT_APPLICABLE = "not_applicable"
SEVERITY_OPTION_MINOR = "minor"
SEVERITY_OPTION_MAJOR = "major"


def _normalize_status_option(value: Any) -> str:
    normalized = _to_text(value, "").strip().lower()
    if not normalized:
        return STATUS_OPTION_FAULT_CLEARED
    if normalized in {
        STATUS_OPTION_FAULT_CLEARED,
        STATUS_OPTION_TEMPORARILY_FIXED,
        STATUS_OPTION_FOLLOW_UP_ACTION_REQUIRED,
    }:
        return normalized
    if "follow" in normalized or "后续" in normalized or "跟进" in normalized:
        return STATUS_OPTION_FOLLOW_UP_ACTION_REQUIRED
    if "temporar" in normalized or "临时" in normalized:
        return STATUS_OPTION_TEMPORARILY_FIXED
    if "clear" in normalized or "cleared" in normalized or "已清除" in normalized:
        return STATUS_OPTION_FAULT_CLEARED
    return STATUS_OPTION_FAULT_CLEARED


def _normalize_severity_option(value: Any) -> str:
    normalized = _to_text(value, "").strip().lower()
    if not normalized:
        return SEVERITY_OPTION_NOT_APPLICABLE
    if normalized in {
        SEVERITY_OPTION_NOT_APPLICABLE,
        SEVERITY_OPTION_MINOR,
        SEVERITY_OPTION_MAJOR,
    }:
        return normalized
    if (
        "major" in normalized
        or "high" in normalized
        or "critical" in normalized
        or "严重" in normalized
        or "重大" in normalized
    ):
        return SEVERITY_OPTION_MAJOR
    if "minor" in normalized or "low" in normalized or "轻微" in normalized:
        return SEVERITY_OPTION_MINOR
    if "not applicable" in normalized or normalized in {"n/a", "na", "不适用"}:
        return SEVERITY_OPTION_NOT_APPLICABLE
    return SEVERITY_OPTION_NOT_APPLICABLE


def _is_placeholder(value: Any) -> bool:
    normalized = _to_text(value, "").strip().lower()
    return normalized in {"", "n/a", "na", "none", "null", "-"}


def _has_valid_event_sequence(events: Any) -> bool:
    if not isinstance(events, list) or not events:
        return False
    for event in events:
        if not isinstance(event, dict):
            continue
        if not _is_placeholder(event.get("event")):
            return True
    return False


def _has_valid_actions(actions: Any) -> bool:
    if not isinstance(actions, list) or not actions:
        return False
    for action in actions:
        if not isinstance(action, dict):
            continue
        if not _is_placeholder(action.get("action")):
            return True
    return False


def validate_required_sections(data: dict[str, Any]) -> list[str]:
    missing_sections = []

    if _is_placeholder(data.get("detailed_description")):
        missing_sections.append("1. Description of the Incident")
    if any(_is_placeholder(data.get(field)) for field in ("start_time", "detection_time", "resolution_time")):
        missing_sections.append("2. Affected Date")
    if not _has_valid_event_sequence(data.get("event_sequence", [])):
        missing_sections.append("3. Event Sequence")

    impact = data.get("impact", {})
    if not isinstance(impact, dict) or _is_placeholder(impact.get("systems")) or _is_placeholder(
        impact.get("severity")
    ):
        missing_sections.append("4. Impact")

    if _is_placeholder(data.get("trigger")) or _is_placeholder(data.get("root_cause")):
        missing_sections.append("5. Root Cause")

    if not _has_valid_actions(data.get("immediate_actions")) and not _has_valid_actions(
        data.get("preventive_actions")
    ):
        missing_sections.append("6. Follow-Up Actions")
    return missing_sections


def normalize_incident_data(raw_data: Any) -> dict[str, Any]:
    source = dict(raw_data) if isinstance(raw_data, dict) else {"detailed_description": _to_text(raw_data)}

    impact_raw = source.get("impact", {})
    impact_map = impact_raw if isinstance(impact_raw, dict) else {}
    body_raw = source.get("report_body", {})
    body_map = body_raw if isinstance(body_raw, dict) else {}
    root_cause_raw = source.get("root_cause")
    root_cause_map = root_cause_raw if isinstance(root_cause_raw, dict) else {}

    start_time = _remove_iso_t_separator(_first_non_empty(source.get("start_time"), default="N/A"))
    detection_time = _remove_iso_t_separator(_first_non_empty(source.get("detection_time"), default=start_time))
    resolution_time = _remove_iso_t_separator(_first_non_empty(source.get("resolution_time"), default=detection_time))

    fault_date, fault_time = _split_date_time(
        _remove_iso_t_separator(_first_non_empty(source.get("fault_date"), start_time, default="N/A"))
    )

    event_sequence = _normalize_event_sequence(source.get("event_sequence"))
    immediate_actions = _normalize_actions(source.get("immediate_actions"), "immediate")
    preventive_actions = _normalize_actions(source.get("preventive_actions"), "preventive")
    status_option = _normalize_status_option(source.get("status_option", source.get("status")))
    severity_option = _normalize_severity_option(source.get("severity_option", source.get("severity")))

    appendix = source.get("appendix", {})
    if not isinstance(appendix, dict):
        appendix = {"notes": _to_text(appendix), "images": []}
    appendix_images = appendix.get("images")
    if not isinstance(appendix_images, list):
        appendix_images = []

    normalized_data = {
        "reference_no": _first_non_empty(source.get("reference_no"), default=_build_reference_no()),
        "fault_date": fault_date,
        "fault_time": _first_non_empty(source.get("fault_time"), default=fault_time),
        "key_facts": {
            "system": _first_non_empty(
                source.get("key_facts", {}).get("system")
                if isinstance(source.get("key_facts"), dict)
                else None,
                source.get("system"),
                default="",
            ),
            "detection_method": _first_non_empty(
                source.get("key_facts", {}).get("detection_method")
                if isinstance(source.get("key_facts"), dict)
                else None,
                source.get("detection_method"),
                default="",
            ),
            "symptoms": _first_non_empty(
                source.get("key_facts", {}).get("symptoms")
                if isinstance(source.get("key_facts"), dict)
                else None,
                source.get("fault_details"),
                source.get("detailed_description"),
                default="",
            ),
        },
        "reporting_person": _first_non_empty(source.get("reporting_person"), default="N/A"),
        "verified_by": _first_non_empty(source.get("verified_by"), default=""),
        "site_id": _first_non_empty(source.get("site_id"), impact_map.get("region"), default="N/A"),
        "system": _first_non_empty(source.get("system"), impact_map.get("systems"), default="N/A"),
        "location": _first_non_empty(source.get("location"), source.get("site_id"), impact_map.get("region"), default="N/A"),
        "fault_details": _first_non_empty(
            source.get("fault_details"),
            source.get("key_facts", {}).get("symptoms") if isinstance(source.get("key_facts"), dict) else None,
            source.get("detailed_description"),
            default="",
        ),
        "arrival_datetime": _remove_iso_t_separator(_first_non_empty(source.get("arrival_datetime"), start_time)),
        "clearance_datetime": _remove_iso_t_separator(_first_non_empty(source.get("clearance_datetime"), resolution_time)),
        "service_person": _first_non_empty(source.get("service_person"), default=""),
        "fault_cause": _first_non_empty(source.get("fault_cause"), default=""),
        "materials_used": _first_non_empty(source.get("materials_used"), default=""),
        "repair_details": _first_non_empty(source.get("repair_details"), default=""),
        "contractor_staff": _first_non_empty(source.get("contractor_staff"), default=""),
        "contractor_signature": _first_non_empty(source.get("contractor_signature"), default=""),
        "contractor_date": _remove_iso_t_separator(_first_non_empty(source.get("contractor_date"), source.get("fault_date"), default=fault_date)),
        "status_option": status_option,
        "status_ref_no": _first_non_empty(source.get("status_ref_no"), default=""),
        "status": _first_non_empty(source.get("status"), default="Fault has been Cleared"),
        "severity_option": severity_option,
        "severity": _first_non_empty(source.get("severity"), impact_map.get("severity"), default=""),
        "comments": _first_non_empty(source.get("comments"), default=""),
        "employer_rep": _first_non_empty(source.get("employer_rep"), default=""),
        "employer_signature": _first_non_empty(source.get("employer_signature"), default=""),
        "closeout_date": _remove_iso_t_separator(_first_non_empty(source.get("closeout_date"), source.get("fault_date"), default=fault_date)),
        "detailed_description": _first_non_empty(
            source.get("detailed_description"), body_map.get("description"), source.get("fault_details"), default=""
        ),
        "affected_date_summary": _remove_iso_t_separator(_first_non_empty(
            source.get("affected_date_summary"),
            body_map.get("affected_date_summary"),
            default=f"{start_time} - {resolution_time}",
        )),
        "start_time": start_time,
        "detection_time": detection_time,
        "resolution_time": resolution_time,
        "total_duration": _first_non_empty(source.get("total_duration"), default=""),
        "event_sequence": event_sequence,
        "impact": {
            "systems": _first_non_empty(impact_map.get("systems"), body_map.get("impact_scope"), source.get("system"), default=""),
            "users": _first_non_empty(impact_map.get("users"), default=""),
            "region": _first_non_empty(impact_map.get("region"), source.get("site_id"), default=""),
            "severity": _first_non_empty(impact_map.get("severity"), body_map.get("impact_severity"), source.get("severity"), default=""),
            "business_impact": _to_lines(
                impact_map.get("business_impact", body_map.get("business_impact", ""))
            ),
        },
        "trigger": _first_non_empty(
            source.get("trigger"),
            root_cause_map.get("proximate_cause"),
            root_cause_map.get("trigger"),
            body_map.get("trigger"),
            source.get("fault_cause"),
            default="",
        ),
        "root_cause": _first_non_empty(
            root_cause_map.get("technical"),
            root_cause_map.get("root_cause"),
            root_cause_map.get("proximate_cause"),
            source.get("root_cause"),
            body_map.get("root_cause"),
            source.get("fault_cause"),
            default="",
        ),
        "root_cause_evidence": _first_non_empty(
            source.get("root_cause_evidence"),
            root_cause_map.get("evidence"),
            root_cause_map.get("process_gap"),
            source.get("evidence"),
            default="",
        ),
        "immediate_actions": immediate_actions,
        "preventive_actions": preventive_actions,
        "allow_incomplete": bool(source.get("allow_incomplete", False)),
        "appendix": {
            "notes": _first_non_empty(appendix.get("notes"), default=""),
            "images": appendix_images,
        },
        "report_body": {
            "description": _first_non_empty(body_map.get("description"), source.get("detailed_description"), default=""),
            "affected_date_summary": _first_non_empty(
                body_map.get("affected_date_summary"),
                source.get("affected_date_summary"),
                default="",
            ),
            "timeline": body_map.get("timeline", []),
            "impact_scope": _first_non_empty(body_map.get("impact_scope"), impact_map.get("systems"), default=""),
            "impact_severity": _first_non_empty(
                body_map.get("impact_severity"),
                impact_map.get("severity"),
                default="",
            ),
            "root_cause": _first_non_empty(body_map.get("root_cause"), source.get("root_cause"), default=""),
            "follow_up_actions": _to_lines(body_map.get("follow_up_actions")),
            "business_impact": _to_lines(body_map.get("business_impact")),
            "trigger": _first_non_empty(body_map.get("trigger"), source.get("trigger"), default=""),
        },
    }
    return normalized_data


def _set_paragraph_text(paragraph, text: str) -> None:
    text_value = _to_text(text, default="")
    if paragraph.runs:
        paragraph.runs[0].text = text_value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text_value)


def _set_cell_text(cell, text: str) -> None:
    if not cell.paragraphs:
        paragraph = cell.add_paragraph()
        _set_paragraph_text(paragraph, text)
        return
    _set_paragraph_text(cell.paragraphs[0], text)
    for paragraph in cell.paragraphs[1:]:
        _set_paragraph_text(paragraph, "")


def _append_cell_text_block(cell, text: str) -> None:
    """在单元格末尾追加内容，避免覆盖模板中的固定标题与占位线。"""
    text_value = _to_text(text, "").strip()
    if not text_value:
        return

    lines = [line.strip() for line in text_value.splitlines() if line.strip()]
    if not lines:
        lines = [text_value]

    for line in lines:
        paragraph = cell.add_paragraph()
        paragraph.paragraph_format.left_indent = None
        paragraph.paragraph_format.first_line_indent = None
        _set_paragraph_text(paragraph, line)


def _set_status_cell_text(cell, *, status_option: str, status_ref_no: str) -> None:
    if not cell.paragraphs:
        paragraph = cell.add_paragraph()
    else:
        paragraph = cell.paragraphs[0]

    template_run = None
    for run in paragraph.runs:
        if (run.text or "").strip():
            template_run = run
            break
    if template_run is None and paragraph.runs:
        template_run = paragraph.runs[0]

    template_font_name = template_run.font.name if template_run is not None else None
    template_font_size = template_run.font.size if template_run is not None else None
    template_bold = template_run.bold if template_run is not None else None
    template_italic = template_run.italic if template_run is not None else None

    if paragraph.runs:
        for run in paragraph.runs:
            run.text = ""
        first_run = paragraph.runs[0]
    else:
        first_run = paragraph.add_run("")

    checked = "☑"
    unchecked = "☐"
    normalized_status_ref = status_ref_no.strip()
    status_ref_display = (
        "_________________"
        if _is_placeholder(normalized_status_ref)
        else normalized_status_ref
    )

    first_run.text = (
        f"{checked if status_option == STATUS_OPTION_FAULT_CLEARED else unchecked} Fault has been Cleared / "
        f"{checked if status_option == STATUS_OPTION_TEMPORARILY_FIXED else unchecked} Temporarily fixed / "
        f"{checked if status_option == STATUS_OPTION_FOLLOW_UP_ACTION_REQUIRED else unchecked} "
        "Follow up action required (Ref No. "
    )
    if template_font_name:
        first_run.font.name = template_font_name
    if template_font_size is not None:
        first_run.font.size = template_font_size
    first_run.bold = template_bold
    first_run.italic = template_italic

    ref_run = paragraph.add_run(status_ref_display)
    ref_run.underline = True
    if template_font_name:
        ref_run.font.name = template_font_name
    if template_font_size is not None:
        ref_run.font.size = template_font_size
    ref_run.bold = template_bold
    ref_run.italic = template_italic

    end_run = paragraph.add_run(")")
    if template_font_name:
        end_run.font.name = template_font_name
    if template_font_size is not None:
        end_run.font.size = template_font_size
    end_run.bold = template_bold
    end_run.italic = template_italic

    for extra_paragraph in cell.paragraphs[1:]:
        _set_paragraph_text(extra_paragraph, "")


def _decode_data_url_to_bytes(data_url: str) -> bytes | None:
    if not isinstance(data_url, str):
        return None
    if not data_url.startswith("data:"):
        return None
    cache_key = hashlib.sha256(data_url.encode("utf-8")).hexdigest()
    cached = _IMAGE_DECODE_CACHE.get(cache_key)
    if cached is not None:
        if cache_key in _IMAGE_DECODE_CACHE_ORDER:
            _IMAGE_DECODE_CACHE_ORDER.remove(cache_key)
        _IMAGE_DECODE_CACHE_ORDER.append(cache_key)
        return cached

    marker = ";base64,"
    index = data_url.find(marker)
    if index < 0:
        return None
    encoded = data_url[index + len(marker) :]
    try:
        decoded = base64.b64decode(encoded)
    except Exception:
        return None
    if len(decoded) > _IMAGE_DECODE_CACHE_MAX_BYTES:
        return decoded

    _IMAGE_DECODE_CACHE[cache_key] = decoded
    if cache_key in _IMAGE_DECODE_CACHE_ORDER:
        _IMAGE_DECODE_CACHE_ORDER.remove(cache_key)
    _IMAGE_DECODE_CACHE_ORDER.append(cache_key)
    while len(_IMAGE_DECODE_CACHE_ORDER) > _IMAGE_DECODE_CACHE_MAX_ENTRIES:
        stale_key = _IMAGE_DECODE_CACHE_ORDER.pop(0)
        _IMAGE_DECODE_CACHE.pop(stale_key, None)
    return decoded


def _sanitize_appendix_notes(notes: Any) -> str:
    raw = _to_text(notes, "").strip()
    if not raw:
        return ""
    text_value = re.sub(r"(?i)<img[^>]*>", "", raw)
    text_value = re.sub(r"(?i)<br\s*/?>", "\n", text_value)
    text_value = re.sub(r"(?i)</p\s*>", "\n", text_value)
    text_value = re.sub(r"(?i)</div\s*>", "\n", text_value)
    text_value = re.sub(r"</?[A-Za-z][^>]*>", "", text_value)
    text_value = re.sub(r"data:image/[^;]+;base64,[A-Za-z0-9+/=]+", "", text_value)
    lines = [unescape(line).strip() for line in text_value.splitlines() if unescape(line).strip()]
    return "\n".join(lines)


def _remove_iso_t_separator(value: str) -> str:
    result = re.sub(r"(\d{4}-\d{2}-\d{2})T(\d{2}:\d{2})", r"\1 \2", value)
    result = re.sub(r"(\d{2}:\d{2}):\d{2}", r"\1", result)
    result = re.sub(r"(\d{2}:\d{2})[+-]\d{2}:\d{2}", r"\1", result)
    result = re.sub(r"(\d{2}:\d{2})Z", r"\1", result)
    return result


def _extract_time_only(datetime_text: str) -> str:
    text = datetime_text.strip()
    matched = re.search(r"(\d{1,2}:\d{2})\s*$", text)
    if matched:
        return matched.group(1)
    return text


class FaultLogFormGenerator:
    """按参考模板填充并输出文档。"""

    def __init__(self, template_path: str | Path | None = None):
        resolved_template = Path(template_path) if template_path else REFERENCE_TEMPLATE_PATH
        if not resolved_template.is_file():
            raise FileNotFoundError(f"参考模板不存在：{resolved_template}")
        self.doc = Document(str(resolved_template))
        # 新逻辑仅支持最新版 incident-report 参考模板。
        if len(self.doc.tables) != 1 or len(self.doc.paragraphs) > 24:
            raise RuntimeError("参考模板结构不匹配：当前仅支持新版 incident-report 模板。")

    def _set_paragraph_if_exists(self, index: int, text: str) -> None:
        if index < 0 or index >= len(self.doc.paragraphs):
            return
        _set_paragraph_text(self.doc.paragraphs[index], text)

    def _find_paragraph_index_by_prefix(self, prefix: str) -> int | None:
        target = prefix.strip().lower()
        if not target:
            return None
        for index, paragraph in enumerate(self.doc.paragraphs):
            text = _to_text(paragraph.text, "").strip().lower()
            if text.startswith(target):
                return index
        return None

    def _set_section_text_after_heading(self, heading_prefix: str, text: str):
        heading_index = self._find_paragraph_index_by_prefix(heading_prefix)
        if heading_index is None:
            return None
        content_index = heading_index + 1
        if content_index < len(self.doc.paragraphs):
            content_para = self.doc.paragraphs[content_index]
            _set_paragraph_text(content_para, text)
            return content_para
        return None

    def _add_blank_line_after_paragraph(self, paragraph) -> None:
        new_p = OxmlElement("w:p")
        paragraph._element.addnext(new_p)

    def _fill_page_one(self, data: dict[str, Any]) -> None:
        table = self.doc.tables[0]
        rows = table.rows

        _set_cell_text(rows[0].cells[1], data.get("reference_no", ""))
        _set_cell_text(rows[2].cells[1], data.get("fault_date", ""))
        _set_cell_text(rows[2].cells[3], data.get("fault_time", ""))
        _set_cell_text(rows[3].cells[1], data.get("reporting_person", ""))
        _set_cell_text(rows[3].cells[3], data.get("verified_by", ""))
        _set_cell_text(rows[4].cells[1], data.get("site_id", ""))
        _set_cell_text(rows[4].cells[3], data.get("system", ""))
        _set_cell_text(rows[5].cells[1], data.get("location", ""))

        detail_text = _to_text(data.get("fault_details"), "").strip() or _to_text(
            data.get("detailed_description"), ""
        ).strip()
        _set_cell_text(rows[6].cells[1], detail_text)

        _set_cell_text(rows[8].cells[1], data.get("arrival_datetime", ""))
        _set_cell_text(rows[8].cells[3], data.get("clearance_datetime", ""))
        _set_cell_text(rows[9].cells[1], data.get("service_person", ""))
        _set_cell_text(rows[9].cells[3], data.get("fault_cause", ""))
        _set_cell_text(rows[10].cells[1], data.get("materials_used", ""))

        _set_cell_text(rows[11].cells[1], data.get("repair_details", ""))

        _set_cell_text(rows[12].cells[1], data.get("contractor_staff", ""))
        _set_cell_text(rows[13].cells[1], data.get("contractor_signature", ""))
        _set_cell_text(rows[14].cells[1], data.get("contractor_date", ""))

        status_option = _normalize_status_option(data.get("status_option", data.get("status")))
        status_ref_no = _to_text(data.get("status_ref_no"), "").strip()
        _set_status_cell_text(
            rows[16].cells[1],
            status_option=status_option,
            status_ref_no=status_ref_no,
        )

        severity_option = _normalize_severity_option(data.get("severity_option", data.get("severity")))
        _set_cell_text(
            rows[17].cells[1],
            f"{'☑' if severity_option == SEVERITY_OPTION_NOT_APPLICABLE else '☐'} Not Applicable",
        )
        _set_cell_text(
            rows[17].cells[2],
            f"{'☑' if severity_option == SEVERITY_OPTION_MINOR else '☐'} Minor",
        )
        _set_cell_text(
            rows[17].cells[3],
            f"{'☑' if severity_option == SEVERITY_OPTION_MAJOR else '☐'} Major",
        )

        _set_cell_text(rows[18].cells[1], data.get("comments", ""))
        _set_cell_text(rows[19].cells[1], data.get("employer_rep", ""))
        _set_cell_text(rows[20].cells[1], data.get("employer_signature", ""))
        _set_cell_text(rows[21].cells[1], data.get("closeout_date", ""))

    def _fill_body_sections(self, data: dict[str, Any]) -> None:
        impact = data.get("impact", {}) if isinstance(data.get("impact"), dict) else {}
        report_body = data.get("report_body", {}) if isinstance(data.get("report_body"), dict) else {}

        description_text = _to_text(data.get("detailed_description"), "")
        content_para = self._set_section_text_after_heading("Description of the Incident:", description_text)
        if content_para is not None:
            self._add_blank_line_after_paragraph(content_para)

        affected_summary = _remove_iso_t_separator(_to_text(data.get("affected_date_summary"), ""))
        timeline_lines: list[str] = []
        events = data.get("event_sequence", [])
        if isinstance(events, list):
            for event in events:
                if not isinstance(event, dict):
                    continue
                time_text = _to_text(event.get("time"), "").strip()
                event_text = _to_text(event.get("event"), "").strip()
                if not time_text and not event_text:
                    continue
                display_time = _extract_time_only(time_text) if time_text else ""
                timeline_lines.append(
                    f"{display_time} - {event_text}".strip(" -")
                    if display_time and event_text
                    else (event_text or display_time)
                )
        affected_lines = [affected_summary] if affected_summary else []
        if timeline_lines:
            affected_lines.append("Timeline:")
            affected_lines.extend(timeline_lines)
        content_para = self._set_section_text_after_heading("Affected Date:", "\n".join(affected_lines))
        if content_para is not None:
            self._add_blank_line_after_paragraph(content_para)

        impact_lines: list[str] = []
        impact_scope = _to_text(impact.get("systems"), "").strip()
        if impact_scope:
            impact_lines.append(f"Scope: {impact_scope}")
        impact_severity = _to_text(impact.get("severity"), "").strip()
        if impact_severity:
            impact_lines.append(f"Severity: {impact_severity}")
        business_impact_lines = _to_lines(impact.get("business_impact"))
        if business_impact_lines:
            impact_lines.append("Business Impact:")
            impact_lines.extend([f"• {line}" for line in business_impact_lines if line.strip()])
        content_para = self._set_section_text_after_heading("Impact:", "\n".join(impact_lines))
        if content_para is not None:
            self._add_blank_line_after_paragraph(content_para)

        root_lines: list[str] = []
        trigger_text = _to_text(data.get("trigger"), "").strip()
        if trigger_text:
            root_lines.append(f"Trigger: {trigger_text}")
        root_cause_text = _to_text(data.get("root_cause"), "").strip()
        if root_cause_text:
            root_lines.append(f"Root Cause: {root_cause_text}")
        root_evidence = _to_text(data.get("root_cause_evidence"), "").strip()
        if root_evidence:
            root_lines.append(f"Evidence: {root_evidence}")
        content_para = self._set_section_text_after_heading("Root Cause:", "\n".join(root_lines))
        if content_para is not None:
            self._add_blank_line_after_paragraph(content_para)

        follow_lines = _to_lines(report_body.get("follow_up_actions"))
        if not follow_lines:
            follow_lines = _extract_action_lines(data.get("preventive_actions"))
        self._set_section_text_after_heading(
            "Follow-Up Actions:",
            "\n".join([f"• {line}" for line in follow_lines if line.strip()]),
        )

    def _fill_appendix(self, data: dict[str, Any]) -> None:
        appendix = data.get("appendix", {})
        if not isinstance(appendix, dict):
            appendix = {}
        notes = _sanitize_appendix_notes(appendix.get("notes"))
        images = appendix.get("images")
        if not isinstance(images, list):
            images = []

        self._set_section_text_after_heading("Appendix:", notes)

        for image in images:
            if not isinstance(image, dict):
                continue
            image_bytes = _decode_data_url_to_bytes(_to_text(image.get("data_url"), ""))
            if image_bytes is None:
                continue
            try:
                image_paragraph = self.doc.add_paragraph("")
                image_run = image_paragraph.add_run()
                image_run.add_picture(io.BytesIO(image_bytes), width=Cm(15))
            except Exception:
                # 个别格式（如部分 webp/svg 转码失败）不应阻断整份事故报告预览。
                continue

    def generate_form(self, data: dict[str, Any], output_path: str | Path | None = None):
        self._fill_page_one(data)
        self._fill_body_sections(data)
        self._fill_appendix(data)

        if output_path:
            self.doc.save(str(output_path))
            print(f"[OK] Fault Log Form saved to: {output_path}")
        return self.doc


def main() -> None:
    import argparse

    parser = argparse.ArgumentParser(description="Generate DAS Fault Log Form")
    parser.add_argument("--json", required=True, type=str, help="Load data from JSON file")
    parser.add_argument("--output", "-o", type=str, default="Fault_Log_Form.docx", help="Output file path")
    args = parser.parse_args()

    print(f"Loading data from {args.json}...")
    with open(args.json, "r", encoding="utf-8") as file:
        data = json.load(file)

    normalized = normalize_incident_data(data)
    if not normalized.get("allow_incomplete"):
        missing_sections = validate_required_sections(normalized)
        if missing_sections:
            missing_text = "; ".join(missing_sections)
            raise SystemExit(
                f"报告信息不完整，缺少以下章节：{missing_text}。"
                "请先补齐正文必填信息；如用户明确接受缺省项，可在 report_data 中设置 allow_incomplete=true。"
            )

    generator = FaultLogFormGenerator()
    generator.generate_form(normalized, args.output)
    print(f"\n[OK] Fault Log Form generated: {args.output}")


if __name__ == "__main__":
    main()
