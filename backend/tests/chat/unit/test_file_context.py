from io import BytesIO
from types import SimpleNamespace

from docx import Document
from openpyxl import Workbook

from doc_process_studio.chat.infrastructure.file_context import (
    decode_file_bytes,
    extract_docx_text,
    extract_pdf_text,
    extract_xlsx_text,
)


def test_decode_file_bytes_supports_utf8_text() -> None:
    assert decode_file_bytes("你好，文档".encode()) == "你好，文档"


def test_extract_docx_text_reads_paragraphs() -> None:
    document = Document()
    document.add_paragraph("这是第一段")
    document.add_paragraph("这是第二段")
    buffer = BytesIO()
    document.save(buffer)

    extracted_text = extract_docx_text(buffer.getvalue())

    assert "这是第一段" in extracted_text
    assert "这是第二段" in extracted_text


def test_extract_xlsx_text_reads_sheet_cells() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "摘要"
    sheet["A1"] = "标题"
    sheet["B1"] = "内容"
    sheet["A2"] = "结论"
    sheet["B2"] = "通过"
    buffer = BytesIO()
    workbook.save(buffer)

    extracted_text = extract_xlsx_text(buffer.getvalue())

    assert "[工作表：摘要]" in extracted_text
    assert "标题 | 内容" in extracted_text
    assert "结论 | 通过" in extracted_text


def test_extract_pdf_text_reads_page_text(monkeypatch) -> None:
    fake_page = SimpleNamespace(extract_text=lambda: "PDF 正文")
    fake_reader = SimpleNamespace(pages=[fake_page])

    monkeypatch.setattr(
        "doc_process_studio.chat.infrastructure.file_context.PdfReader",
        lambda _: fake_reader,
    )

    extracted_text = extract_pdf_text(b"%PDF-FAKE")

    assert extracted_text == "PDF 正文"
