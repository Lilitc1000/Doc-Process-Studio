import json
from pathlib import Path
from typing import Any

import pytest
from docx.api import Document

from doc_process_studio.chat.infrastructure import attachments as attachments_module
from doc_process_studio.chat.router.schemas.request import ChatStreamRequest
from doc_process_studio.skill.application.dtos.runtime import SkillConversationState
from doc_process_studio.skill.infrastructure.tool_loop import execute_skill_tool_call


def _build_tool_call(arguments: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": "tool-call-incident-report",
        "type": "function",
        "function": {
            "name": "generate_incident_report",
            "arguments": json.dumps(arguments, ensure_ascii=False),
        },
    }


def _read_cell(table: Any, row_index: int, cell_index: int) -> str:
    return str(table.rows[row_index].cells[cell_index].text.strip())


def _find_paragraph_after_heading(document: Any, heading_prefix: str) -> str:
    target = heading_prefix.strip().lower()
    for index, paragraph in enumerate(document.paragraphs):
        if paragraph.text.strip().lower().startswith(target):
            if index + 1 < len(document.paragraphs):
                return str(document.paragraphs[index + 1].text.strip())
            return ""
    return ""


async def test_incident_report_tool_chain_generates_non_empty_key_cells(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        attachments_module.settings,
        "generated_attachments_dir",
        str(tmp_path),
    )

    request = ChatStreamRequest(
        user_message_id="msg-user-incident-1",
        conversation_id="conv-incident-1",
        model="qwen3-coder-next:latest",
        selected_skill_ids=["incident-report"],
        messages=[],
        attachment_ids=[],
    )
    state = SkillConversationState(
        conversation_id=request.conversation_id,
        skill_id="incident-report",
        system_prompt="",
    )

    # 故意传入"半结构化"数据，验证真实 tools.json 调用链下脚本仍能补齐模板关键字段。
    report_data = {
        "reference_no": "DAS-20260408-777",
        "detailed_description": "Payment service outage",
        "key_facts": {
            "system": "Payment Gateway",
            "detection_method": "Monitoring alert",
            "symptoms": "Transaction timeout",
        },
        "start_time": "08/04/2026 09:10",
        "detection_time": "08/04/2026 09:12",
        "resolution_time": "08/04/2026 09:40",
        "severity": "P2",
        "impact": {
            "systems": "Payment Gateway",
            "severity": "High",
            "business_impact": ["Users cannot complete payment transactions"],
        },
        "event_sequence": ["09:12 Monitoring alert triggered"],
        "root_cause": {
            "technical": "Missing DB index caused full table scan",
            "proximate_cause": "SQL deployed without pre-release benchmark",
            "process_gap": "Release checklist missing DBA review",
        },
        "immediate_actions": ["Restarted payment service"],
        "preventive_actions": ["Add payment timeout alert rule"],
        "allow_incomplete": True,
    }
    tool_call = _build_tool_call(
        {
            "report_data": report_data,
            "output_name": "incident-report-integration.docx",
        }
    )

    tool_result, attachments = await execute_skill_tool_call(
        request=request,
        state=state,
        tool_call=tool_call,
    )

    assert tool_result.get("ok") is True
    assert len(attachments) == 1

    attachment = attachments[0]
    metadata, attachment_path, is_expired = attachments_module.resolve_attachment_path(attachment.attachment_id)

    assert metadata is not None
    assert is_expired is False
    assert attachment_path is not None
    assert attachment_path.suffix.lower() == ".docx"
    assert attachment_path.is_file()

    document = Document(str(attachment_path))
    # 新参考模板为"1 张首页表格 + 正文段落"结构。
    assert len(document.tables) == 1

    page_one_table = document.tables[0]

    # 断言 Page 1 关键字段单元格不是空字符串，避免"模板区域空白"回归。
    assert _read_cell(page_one_table, 0, 1) != ""  # Reference No.
    assert _read_cell(page_one_table, 4, 1) != ""  # Site ID
    assert _read_cell(page_one_table, 5, 1) != ""  # Location of Fault
    assert _read_cell(page_one_table, 6, 1) != ""  # Details of Fault Symptom 内容
    assert _read_cell(page_one_table, 16, 1) != ""  # Status
    assert "Details of Fault" in _read_cell(page_one_table, 6, 0)
    assert "Symptom:" in _read_cell(page_one_table, 6, 0)
    assert "Details of repair works" in _read_cell(page_one_table, 11, 0)
    assert "Ref No." in _read_cell(page_one_table, 16, 1)

    # 正文段落应写入核心章节内容。
    assert _find_paragraph_after_heading(document, "Description of the Incident:") != ""
    assert _find_paragraph_after_heading(document, "Affected Date:") != ""
    assert _find_paragraph_after_heading(document, "Impact:") != ""
    assert _find_paragraph_after_heading(document, "Root Cause:") != ""
    assert _find_paragraph_after_heading(document, "Follow-Up Actions:") != ""

    all_paragraph_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert '{"process_gap"' not in all_paragraph_text
